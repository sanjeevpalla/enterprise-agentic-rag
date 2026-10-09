"""Microsoft Office loaders: Word (.docx), Excel (.xlsx) and PowerPoint (.pptx)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import docx
import openpyxl
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.loaders.base import BaseLoader, Document, LoaderError


class DocxLoader(BaseLoader):
    """Load paragraphs and tables from Word documents, in document order.

    Args:
        split_by_headings: If True (default), return one Document per heading section
            (Word's built-in "Title" / "Heading N" styles). Each records its heading path as
            ``h1``/``h2``/``h3`` metadata (deeper headings: ``section_heading`` only), which
            citations show and the splitter prepends to every chunk; the heading text isn't
            repeated in the content.
        include_tables: Include table contents as pipe-separated rows.
    """

    supported_extensions = (".docx",)

    def __init__(self, split_by_headings: bool = True, include_tables: bool = True) -> None:
        self.split_by_headings = split_by_headings
        self.include_tables = include_tables

    def _load(self, path: Path) -> list[Document]:
        document = docx.Document(str(path))
        base_metadata = self._extract_metadata(document, path)

        # (heading metadata, paragraphs) per section; the first holds text before any heading.
        sections: list[tuple[dict[str, str], list[str]]] = [({}, [])]
        path_by_level: dict[int, str] = {}
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                level = self._heading_level(block) if self.split_by_headings else None
                if level is not None:
                    path_by_level = {lvl: t for lvl, t in path_by_level.items() if lvl < level}
                    path_by_level[level] = text
                    headings = {f"h{lvl}": t for lvl, t in path_by_level.items() if lvl <= 3}
                    sections.append(({**headings, "section_heading": text}, []))
                else:
                    sections[-1][1].append(text)
            elif isinstance(block, Table) and self.include_tables:
                table_text = self._table_text(block)
                if table_text:
                    sections[-1][1].append(table_text)

        documents: list[Document] = []
        for index, (headings, parts) in enumerate(sections):
            content = "\n\n".join(parts)
            if not content.strip():
                continue
            metadata = dict(base_metadata)
            if self.split_by_headings:
                metadata.update(headings, section_index=index)
            documents.append(Document(page_content=content, metadata=metadata))

        if not documents:
            raise LoaderError(f"No extractable text found in {path}")
        return documents

    @staticmethod
    def _heading_level(paragraph: Paragraph) -> int | None:
        """1 for "Title", N for "Heading N", None for body text."""
        style = paragraph.style.name if paragraph.style is not None else ""
        if style == "Title":
            return 1
        if style.startswith("Heading"):
            level = style.removeprefix("Heading").strip()
            return int(level) if level.isdigit() else 1
        return None

    @staticmethod
    def _table_text(table: Table) -> str:
        rows = []
        for row in table.rows:
            # Merged cells repeat across the row; drop consecutive duplicates.
            cells: list[str] = []
            for cell in row.cells:
                text = cell.text.strip()
                if not cells or cells[-1] != text:
                    cells.append(text)
            if any(cells):
                rows.append(" | ".join(cells))
        return "\n".join(rows)

    @staticmethod
    def _extract_metadata(document: Any, path: Path) -> dict[str, Any]:
        props = document.core_properties
        metadata: dict[str, Any] = {"source": str(path), "file_type": "docx"}
        if props.title:
            metadata["title"] = props.title
        if props.author:
            metadata["author"] = props.author
        if props.created:
            metadata["created"] = props.created.isoformat()
        return metadata


class XlsxLoader(BaseLoader):
    """Load Excel workbooks, one Document per sheet.

    Args:
        header_row: Treat the first non-empty row as column headers and render
            each following row as ``header: value`` pairs, which keeps column
            meaning attached to every row for retrieval.
    """

    supported_extensions = (".xlsx", ".xlsm")

    def __init__(self, header_row: bool = True) -> None:
        self.header_row = header_row

    def _load(self, path: Path) -> list[Document]:
        # data_only=True returns cached formula results instead of the formulas.
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            documents = [
                document
                for sheet in workbook.worksheets
                if (document := self._sheet_document(sheet, path)) is not None
            ]
        finally:
            workbook.close()

        if not documents:
            raise LoaderError(f"No data found in {path}")
        return documents

    def _sheet_document(self, sheet: Any, path: Path) -> Document | None:
        rows = [
            ["" if value is None else str(value).strip() for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(row)]
        if not rows:
            return None

        if self.header_row and len(rows) > 1:
            headers = rows[0]
            lines = [
                ", ".join(
                    f"{headers[i] or f'col{i + 1}'}: {value}"
                    for i, value in enumerate(row)
                    if value and i < len(headers)
                )
                for row in rows[1:]
            ]
        else:
            lines = [" | ".join(row) for row in rows]

        metadata = {
            "source": str(path),
            "file_type": "xlsx",
            "sheet": sheet.title,
            "row_count": len(rows),
        }
        return Document(page_content=f"Sheet: {sheet.title}\n" + "\n".join(lines), metadata=metadata)


class PptxLoader(BaseLoader):
    """Load PowerPoint decks, one Document per slide.

    Args:
        include_notes: Append the slide's speaker notes.
    """

    supported_extensions = (".pptx",)

    def __init__(self, include_notes: bool = True) -> None:
        self.include_notes = include_notes

    def _load(self, path: Path) -> list[Document]:
        presentation = pptx.Presentation(str(path))
        documents: list[Document] = []

        for number, slide in enumerate(presentation.slides, start=1):
            title = slide.shapes.title.text.strip() if slide.shapes.title else None
            parts = [text for shape in slide.shapes if (text := self._shape_text(shape))]

            if self.include_notes and slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"Notes: {notes}")

            if not parts:
                continue
            metadata = {
                "source": str(path),
                "file_type": "pptx",
                "slide": number,
                "slide_title": title,
                "total_slides": len(presentation.slides),
            }
            documents.append(Document(page_content="\n\n".join(parts), metadata=metadata))

        if not documents:
            raise LoaderError(f"No extractable text found in {path}")
        return documents

    def _shape_text(self, shape: Any) -> str:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            return "\n".join(t for s in shape.shapes if (t := self._shape_text(s)))
        if shape.has_text_frame:
            return shape.text_frame.text.strip()
        if shape.has_table:
            rows = ([cell.text.strip() for cell in row.cells] for row in shape.table.rows)
            return "\n".join(" | ".join(cells) for cells in rows if any(cells))
        return ""


class OfficeLoader(BaseLoader):
    """Dispatch to the right Office loader based on file extension.

    Pass pre-configured loaders to override the defaults.
    """

    def __init__(
        self,
        docx_loader: DocxLoader | None = None,
        xlsx_loader: XlsxLoader | None = None,
        pptx_loader: PptxLoader | None = None,
    ) -> None:
        loaders: list[BaseLoader] = [
            docx_loader or DocxLoader(),
            xlsx_loader or XlsxLoader(),
            pptx_loader or PptxLoader(),
        ]
        self._by_extension = {ext: loader for loader in loaders for ext in loader.supported_extensions}
        self.supported_extensions = tuple(self._by_extension)

    def _load(self, path: Path) -> list[Document]:
        return self._by_extension[path.suffix.lower()]._load(path)
